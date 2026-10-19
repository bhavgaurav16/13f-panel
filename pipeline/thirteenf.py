#!/usr/bin/env python3
"""
13F capital-flows tracker.

Downloads SEC Form 13F bulk data sets, rebuilds each filer's long-equity book for
consecutive quarters, and aggregates share-count changes (price-neutral flows)
into AUM tiers. Outputs:
  out/history.json  - all quarters computed so far (merged with any existing history)
  out/13F_flows_<Q>.xlsx - scan-friendly workbook for the latest quarter + trends

Usage:
  python thirteenf.py --quarters 4            # latest 4 quarters of flows
  python thirteenf.py --quarters 1 --history existing_history.json
"""
import argparse, io, json, os, re, sys, zipfile, datetime as dt
from collections import defaultdict
import numpy as np
import pandas as pd
import requests

UA = {"User-Agent": os.environ.get("SEC_USER_AGENT", "13F research contact@example.com"), "Accept-Encoding": "gzip, deflate"}
PAGE = "https://www.sec.gov/data-research/sec-markets-data/form-13f-data-sets"
CACHE = os.environ.get("THIRTEENF_CACHE", os.path.join(os.path.dirname(os.path.abspath(__file__)), "cache"))
OUT = os.environ.get("THIRTEENF_OUT", os.path.join(os.path.dirname(os.path.abspath(__file__)), "out"))

TIERS = [  # (key, label, lower bound in $), top10 assigned first
    ("top10", "Top 10 managers", None),
    ("t10b", "$10B+ (ex top 10)", 10e9),
    ("t1b", "$1B-10B", 1e9),
    ("t100m", "$100M-1B", 100e6),
]
TOPN = 20
ETF_PAT = re.compile(r"\b(?:ETF|ETN|SPDR|ISHARES|INDEX FD|INDEX FUND|SELECT SECTOR|VANGUARD|INVESCO QQQ|PROSHARES|DIREXION|"
                     r"WISDOMTREE|SCHWAB STRATEGIC|TRUST UNIT|SER TR|ETF TR|FDS TR|EXCHANGE TRADED|JPMORGAN EXCHANGE|"
                     r"DIMENSIONAL ETF|FIRST TR|GLOBAL X|ARK ETF|VANECK|PIMCO ETF|AMPLIFY|BLACKROCK ETF|ISHARES TR|SPDR S&P|"
                     r"SPDR SER|SPDR GOLD|SPDR INDEX|SECTOR SPDR)\b", re.I)


# ----------------------------------------------------------------------------- download
def list_zips():
    html = requests.get(PAGE, headers=UA, timeout=60).text
    urls = sorted(set(re.findall(r'href="([^"]+_form13f\.zip)"', html)))
    urls = [u if u.startswith("http") else "https://www.sec.gov" + u for u in urls]
    # keep dated-window files only (older files are named like 2013q2_form13f.zip; panel.py handles those)
    urls = [u for u in urls if re.search(r"\d{2}[a-z]{3}\d{4}-\d{2}[a-z]{3}\d{4}_form13f", u.lower())]

    def start(u):
        m = re.search(r"(\d{2}[a-z]{3}\d{4})-\d{2}[a-z]{3}\d{4}_form13f", u)
        return dt.datetime.strptime(m.group(1), "%d%b%Y")
    return sorted(urls, key=start)


def fetch(url):
    os.makedirs(CACHE, exist_ok=True)
    path = os.path.join(CACHE, url.rsplit("/", 1)[-1])
    if not os.path.exists(path):
        print("downloading", url, file=sys.stderr)
        with requests.get(url, headers=UA, timeout=600, stream=True) as r:
            r.raise_for_status()
            with open(path + ".part", "wb") as f:
                for chunk in r.iter_content(1 << 20):
                    f.write(chunk)
        os.replace(path + ".part", path)
    return path


def read_zip(path):
    z = zipfile.ZipFile(path)
    names = {n.rsplit("/", 1)[-1].upper(): n for n in z.namelist()}

    def rd(name, cols=None):
        with z.open(names[name]) as f:
            df = pd.read_csv(f, sep="\t", dtype=str, usecols=cols, quoting=3, on_bad_lines="skip", low_memory=False)
        df.columns = [c.upper() for c in df.columns]
        return df
    sub = rd("SUBMISSION.TSV")
    cov = rd("COVERPAGE.TSV")
    info = rd("INFOTABLE.TSV")
    return sub, cov, info


# ----------------------------------------------------------------------------- books
def qlabel(d):
    return f"{d.year}Q{(d.month - 1) // 3 + 1}"


def build_books(zip_paths):
    subs, covs, infos = [], [], []
    for p in zip_paths:
        s, c, i = read_zip(p)
        subs.append(s); covs.append(c)
        keep = ["ACCESSION_NUMBER", "NAMEOFISSUER", "TITLEOFCLASS", "CUSIP", "VALUE", "SSHPRNAMT", "SSHPRNAMTTYPE", "PUTCALL"]
        infos.append(i[[k for k in keep if k in i.columns]])
    sub = pd.concat(subs).drop_duplicates("ACCESSION_NUMBER")
    cov = pd.concat(covs).drop_duplicates("ACCESSION_NUMBER")
    info = pd.concat(infos)
    del subs, covs, infos

    sub = sub[sub["SUBMISSIONTYPE"].str.startswith("13F-HR", na=False)]
    sub["PERIOD"] = pd.to_datetime(sub["PERIODOFREPORT"], format="%d-%b-%Y", errors="coerce")
    sub["FILED"] = pd.to_datetime(sub["FILING_DATE"], format="%d-%b-%Y", errors="coerce")
    sub = sub.merge(cov[["ACCESSION_NUMBER", "FILINGMANAGER_NAME", "ISAMENDMENT", "AMENDMENTTYPE"]],
                    on="ACCESSION_NUMBER", how="left")
    sub["Q"] = sub["PERIOD"].map(lambda d: qlabel(d) if pd.notna(d) else None)

    # choose accessions per (CIK, Q): latest restatement-or-original base, plus later NEW HOLDINGS amendments
    sub["ISAMEND"] = (sub["ISAMENDMENT"].fillna("N").str.upper() == "Y") | sub["SUBMISSIONTYPE"].str.endswith("/A")
    sub["NEWH"] = sub["ISAMEND"] & sub["AMENDMENTTYPE"].fillna("").str.upper().str.contains("NEW HOLDINGS")
    base = sub[~sub["NEWH"]].sort_values("FILED").groupby(["CIK", "Q"]).tail(1)
    newh = sub[sub["NEWH"]].merge(base[["CIK", "Q", "FILED"]].rename(columns={"FILED": "BFILED"}), on=["CIK", "Q"])
    newh = newh[newh["FILED"] >= newh["BFILED"]]
    chosen = pd.concat([base, newh[base.columns]])
    acc = chosen[["ACCESSION_NUMBER", "CIK", "Q", "FILINGMANAGER_NAME"]]

    info = info.merge(acc, on="ACCESSION_NUMBER", how="inner")
    info["VALUE"] = pd.to_numeric(info["VALUE"], errors="coerce").fillna(0.0).astype("float64")
    info["SSHPRNAMT"] = pd.to_numeric(info["SSHPRNAMT"], errors="coerce").fillna(0.0).astype("float64")
    info["CUSIP"] = info["CUSIP"].str.upper().str.strip().str[:9]
    info["PUTCALL"] = info["PUTCALL"].fillna("").str.upper().str.strip()
    names = chosen.sort_values("FILED").groupby("CIK")["FILINGMANAGER_NAME"].last()
    return info, names


def fix_units(info):
    """Some filers still report VALUE in $000s. Detect per filer-quarter and rescale."""
    eq = info[(info["SSHPRNAMTTYPE"] == "SH") & (info["PUTCALL"] == "") & (info["SSHPRNAMT"] > 0) & (info["VALUE"] > 0)].copy()
    eq["PX"] = eq["VALUE"] / eq["SSHPRNAMT"]
    med = eq.groupby(["Q", "CUSIP"])["PX"].median().rename("MPX")
    eq = eq.join(med, on=["Q", "CUSIP"])
    eq["R"] = eq["PX"] / eq["MPX"]
    fr = eq.groupby(["CIK", "Q"])["R"].median()
    bad = fr[fr < 0.01].index
    if len(bad):
        idx = info.set_index(["CIK", "Q"]).index.isin(bad)
        info.loc[idx, "VALUE"] *= 1000.0
        print(f"rescaled {len(bad)} filer-quarters reporting in $000s", file=sys.stderr)
    return info


def quarter_tables(info, q):
    d = info[info["Q"] == q]
    long_ = d[(d["SSHPRNAMTTYPE"] == "SH") & (d["PUTCALL"] == "")]
    book = long_.groupby(["CIK", "CUSIP"], as_index=False).agg(SH=("SSHPRNAMT", "sum"), VAL=("VALUE", "sum"))
    issuer = long_.groupby("CUSIP")["NAMEOFISSUER"].agg(lambda s: s.mode().iat[0] if len(s.mode()) else s.iat[0])
    opts = d[d["PUTCALL"].isin(["PUT", "CALL"])].groupby(["CIK", "PUTCALL"])["VALUE"].sum().unstack(fill_value=0)
    return book, issuer, opts


def prev_q(q):
    y, n = int(q[:4]), int(q[-1])
    return f"{y - 1}Q4" if n == 1 else f"{y}Q{n - 1}"


# ----------------------------------------------------------------------------- flows
def compute_quarter(info, names, q):
    pq = prev_q(q)
    cur, issuer_c, opts = quarter_tables(info, q)
    prv, issuer_p, _ = quarter_tables(info, pq)
    issuer = issuer_c.combine_first(issuer_p)

    aum = cur.groupby("CIK")["VAL"].sum()
    aum_p = prv.groupby("CIK")["VAL"].sum()
    both = aum.index.intersection(aum_p.index)
    new_filers = aum.index.difference(aum_p.index)
    gone_filers = aum_p.index.difference(aum.index)

    # tiers on current-quarter long AUM
    tier = pd.Series("sub100m", index=aum.index)
    top10 = aum.sort_values(ascending=False).head(10).index
    for key, _, lo in TIERS[::-1]:
        if lo is not None:
            tier[aum >= lo] = key
    tier[top10] = "top10"

    # price per CUSIP (median implied), current then prior as fallback
    def px(book):
        b = book[(book.SH > 0) & (book.VAL > 0)]
        return (b.VAL / b.SH).groupby(b.CUSIP).median()
    px_c, px_p = px(cur), px(prv)

    m = cur[cur.CIK.isin(both)].merge(prv[prv.CIK.isin(both)], on=["CIK", "CUSIP"], how="outer", suffixes=("", "_P")).fillna(0.0)

    # split detection: median share ratio across continuing holders
    hold = m[(m.SH > 0) & (m.SH_P > 0)]
    ratio = (hold.SH / hold.SH_P).groupby(hold.CUSIP).agg(["median", "count"])
    split = ratio[(ratio["count"] >= 5) & ((ratio["median"] > 1.8) | (ratio["median"] < 0.55))]["median"]
    common = np.array([1/10, 1/8, 1/5, 1/4, 1/3, 1/2, 2, 3, 4, 5, 8, 10, 15, 20, 25, 30, 40, 50])
    split = split.map(lambda r: common[np.argmin(np.abs(np.log(common / r)))] if np.min(np.abs(np.log(common / r))) < 0.08 else np.nan).dropna()
    m["ADJ"] = m.CUSIP.map(split).fillna(1.0)
    m["SH_P"] = m.SH_P * m.ADJ

    m["PX"] = m.CUSIP.map(px_c).fillna(m.CUSIP.map(px_p) / m.ADJ)
    m["PX"] = m.PX.fillna(0.0)
    m["DSH"] = m.SH - m.SH_P
    m["FLOW"] = m.DSH * m.PX
    m.loc[m.DSH.abs() < 1, "FLOW"] = 0.0
    m["TIER"] = m.CIK.map(tier)
    m["NEW"] = (m.SH_P <= 0) & (m.SH > 0)
    m["EXIT"] = (m.SH_P > 0) & (m.SH <= 0)
    m["ETF"] = m.CUSIP.map(issuer).fillna("").str.contains(ETF_PAT)

    out = {"quarter": q, "prior": pq, "generated": dt.date.today().isoformat(),
           "splits_adjusted": {issuer.get(k, k): float(v) for k, v in split.items()},
           "universe": {"filers": int(len(aum)), "continuing": int(len(both)),
                        "new_filers": int(len(new_filers)), "dropped_filers": int(len(gone_filers))},
           "tiers": {}, "managers": []}

    tier_keys = [t[0] for t in TIERS] + ["all"]
    labels = dict((k, l) for k, l, _ in TIERS); labels["all"] = "All 13F filers ($100M+)"
    for tk in tier_keys:
        sel = m[m.TIER.isin([t[0] for t in TIERS])] if tk == "all" else m[m.TIER == tk]
        ciks = sel.CIK.unique()
        a_c = aum.reindex(ciks).sum(); a_p = aum_p.reindex(ciks).sum()
        buys = sel.FLOW.clip(lower=0).sum(); sells = -sel.FLOW.clip(upper=0).sum()
        stk = sel[~sel.ETF]; etf = sel[sel.ETF]
        put = opts.reindex(ciks).get("PUT", pd.Series(dtype=float)).sum() if "PUT" in opts else 0.0
        call = opts.reindex(ciks).get("CALL", pd.Series(dtype=float)).sum() if "CALL" in opts else 0.0
        # concentration: avg share of book in top-10 positions
        bk = cur[cur.CIK.isin(ciks)].sort_values(["CIK", "VAL"], ascending=[True, False])
        top10sh = (bk.groupby("CIK").head(10).groupby("CIK").VAL.sum() / aum.reindex(ciks)).median()
        npos = cur[cur.CIK.isin(ciks)].groupby("CIK").size().median()

        def rank(df, asc):
            g = df.groupby("CUSIP").agg(flow=("FLOW", "sum"), buyers=("FLOW", lambda s: int((s > 0).sum())),
                                        sellers=("FLOW", lambda s: int((s < 0).sum())), new=("NEW", "sum"),
                                        exits=("EXIT", "sum"), held_val=("VAL", "sum"))
            g = g.sort_values("flow", ascending=asc).head(TOPN)
            g = g[(g.flow < 0) if asc else (g.flow > 0)]
            return [{"cusip": c, "issuer": str(issuer.get(c, c)).title(), "flow": float(r.flow),
                     "buyers": int(r.buyers), "sellers": int(r.sellers), "new": int(r.new), "exits": int(r.exits),
                     "held": float(r.held_val), "flow_pct_held": float(r.flow / r.held_val) if r.held_val else None}
                    for c, r in g.iterrows()]

        def breadth(df, asc):
            g = df.groupby("CUSIP").agg(b=("FLOW", lambda s: int((s > 0).sum())), s=("FLOW", lambda s: int((s < 0).sum())),
                                        flow=("FLOW", "sum"))
            g["net"] = g.b - g.s
            g = g.sort_values("net", ascending=asc).head(10)
            return [{"cusip": c, "issuer": str(issuer.get(c, c)).title(), "net_managers": int(r.net), "buyers": int(r.b),
                     "sellers": int(r.s), "flow": float(r.flow)} for c, r in g.iterrows()]

        out["tiers"][tk] = {
            "label": labels[tk], "managers": int(len(ciks)), "aum": float(a_c), "aum_prior": float(a_p),
            "gross_buys": float(buys), "gross_sells": float(sells), "net_flow": float(buys - sells),
            "net_flow_pct": float((buys - sells) / a_p) if a_p else None,
            "turnover": float(min(buys, sells) / ((a_c + a_p) / 2)) if a_c else None,
            "gross_activity": float((buys + sells) / ((a_c + a_p) / 2)) if a_c else None,
            "etf_net_flow": float(etf.FLOW.sum()), "stock_net_flow": float(stk.FLOW.sum()),
            "new_positions": int(sel.NEW.sum()), "full_exits": int(sel.EXIT.sum()),
            "put_notional": float(put), "call_notional": float(call),
            "put_call": float(put / call) if call else None,
            "median_top10_conc": float(top10sh) if pd.notna(top10sh) else None,
            "median_positions": float(npos) if pd.notna(npos) else None,
            "top_buys": rank(stk, False), "top_sells": rank(stk, True),
            "etf_buys": rank(etf, False)[:10], "etf_sells": rank(etf, True)[:10],
            "most_bought_by_count": breadth(stk, False), "most_sold_by_count": breadth(stk, True),
        }

    # top-10 manager detail
    for cik in top10:
        s = m[(m.CIK == cik) & (~m.ETF)]
        s2 = s.groupby("CUSIP").FLOW.sum()
        out["managers"].append({
            "cik": cik, "name": str(names.get(cik, cik)).title(), "aum": float(aum[cik]), "aum_prior": float(aum_p.get(cik, np.nan)),
            "net_flow": float(m[m.CIK == cik].FLOW.sum()),
            "turnover": float(min(m[m.CIK == cik].FLOW.clip(lower=0).sum(), -m[m.CIK == cik].FLOW.clip(upper=0).sum()) / aum[cik]),
            "buys": [{"issuer": str(issuer.get(c, c)).title(), "flow": float(v)} for c, v in s2.nlargest(5).items() if v > 0],
            "sells": [{"issuer": str(issuer.get(c, c)).title(), "flow": float(v)} for c, v in s2.nsmallest(5).items() if v < 0],
        })
    return out, m, aum, tier, names


# ----------------------------------------------------------------------------- excel
def write_excel(hist, latest, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F3A5F")
    T = Font(bold=True, size=14); B = Font(bold=True)
    usd = '#,##0.0,,,"bn";[Red]-#,##0.0,,,"bn"'
    usdm = '#,##0,,"m";[Red]-#,##0,,"m"'
    pct = '0.0%;[Red]-0.0%'
    q = latest["quarter"]

    def header(ws, row, cols):
        for j, c in enumerate(cols, 1):
            x = ws.cell(row, j, c); x.font = H; x.fill = HF; x.alignment = Alignment(wrap_text=True, vertical="center")

    def widths(ws, w):
        for i, v in enumerate(w, 1):
            ws.column_dimensions[get_column_letter(i)].width = v

    # Summary
    ws = wb.active; ws.title = "Summary"
    ws["A1"] = f"13F capital flows - {q} vs {latest['prior']}"; ws["A1"].font = T
    u = latest["universe"]
    ws["A2"] = (f"{u['filers']:,} filers; {u['continuing']:,} filed both quarters (flows use these only). "
                f"Flows = change in share count x current price, so price moves are stripped out. "
                f"Long equity only, so puts and calls are excluded from flows. Stock rankings exclude ETFs.")
    cols = ["Tier", "Managers", "AUM", "Net flow", "Net flow % AUM", "Gross buys", "Gross sells", "Turnover",
            "Stock net", "ETF net", "New positions", "Full exits", "Put/Call notional", "Median top-10 conc.", "Median # positions"]
    header(ws, 4, cols)
    r = 5
    for tk, t in latest["tiers"].items():
        vals = [t["label"], t["managers"], t["aum"], t["net_flow"], t["net_flow_pct"], t["gross_buys"], t["gross_sells"],
                t["turnover"], t["stock_net_flow"], t["etf_net_flow"], t["new_positions"], t["full_exits"], t["put_call"],
                t["median_top10_conc"], t["median_positions"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v)
            if j in (3, 4, 6, 7, 9, 10): c.number_format = usd
            if j in (5, 8, 14): c.number_format = pct
            if j == 13: c.number_format = "0.00"
        if tk == "all":
            for j in range(1, len(cols) + 1): ws.cell(r, j).font = B
        r += 1
    ws.conditional_formatting.add(f"E5:E{r-1}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="max", end_color="63BE7B"))
    widths(ws, [26, 10, 12, 12, 12, 12, 12, 10, 12, 12, 10, 10, 10, 12, 12])
    ws.freeze_panes = "B5"

    # Cross-tier consensus: names in top buys of 3+ tiers
    r += 1; ws.cell(r, 1, "Cross-tier consensus (stock in top 20 of 3+ tiers)").font = T; r += 1
    header(ws, r, ["Issuer", "Direction", "Tiers", "Total flow"]); r += 1
    for side in ("top_buys", "top_sells"):
        cnt, tot, tl = defaultdict(int), defaultdict(float), defaultdict(list)
        for tk in [t[0] for t in TIERS]:
            for x in latest["tiers"][tk][side]:
                cnt[x["issuer"]] += 1; tot[x["issuer"]] += x["flow"]; tl[x["issuer"]].append(tk)
        for n in sorted([k for k in cnt if cnt[k] >= 3], key=lambda k: -abs(tot[k])):
            ws.cell(r, 1, n); ws.cell(r, 2, "BUY" if side == "top_buys" else "SELL"); ws.cell(r, 3, ", ".join(tl[n]))
            ws.cell(r, 4, tot[n]).number_format = usd; r += 1

    # Per tier sheets
    for tk, _, _ in TIERS + [("all", None, None)]:
        t = latest["tiers"][tk]
        ws = wb.create_sheet({"top10": "Top 10", "t10b": "$10B+", "t1b": "$1-10B", "t100m": "$100M-1B", "all": "All"}[tk])
        ws["A1"] = f"{t['label']} - {q}"; ws["A1"].font = T
        ws["A2"] = f"{t['managers']} managers | net {t['net_flow']/1e9:,.1f}bn | turnover {t['turnover']:.1%} | P/C {t['put_call'] or 0:.2f}"
        cols = ["#", "Top buys", "Net $", "Buyers", "Sellers", "New", "Flow % of held", "", "#", "Top sells", "Net $", "Buyers", "Sellers", "Exits", "Flow % of held"]
        header(ws, 4, cols)
        for i in range(TOPN):
            for off, side, ex in ((0, "top_buys", "new"), (8, "top_sells", "exits")):
                if i < len(t[side]):
                    x = t[side][i]
                    row = [i + 1, x["issuer"], x["flow"], x["buyers"], x["sellers"], x[ex], x["flow_pct_held"]]
                    for j, v in enumerate(row, 1):
                        c = ws.cell(5 + i, off + j, v)
                        if j == 3: c.number_format = usdm
                        if j == 7: c.number_format = pct
        r = 6 + TOPN
        ws.cell(r, 1, "Breadth: most managers adding (net count)").font = B
        ws.cell(r, 9, "Breadth: most managers cutting (net count)").font = B
        header(ws, r + 1, ["#", "Issuer", "Net mgrs", "Buyers", "Sellers", "Net $", "", "", "#", "Issuer", "Net mgrs", "Buyers", "Sellers", "Net $"])
        for i in range(10):
            for off, side in ((0, "most_bought_by_count"), (8, "most_sold_by_count")):
                if i < len(t[side]):
                    x = t[side][i]
                    for j, v in enumerate([i + 1, x["issuer"], x["net_managers"], x["buyers"], x["sellers"], x["flow"]], 1):
                        c = ws.cell(r + 2 + i, off + j, v)
                        if j == 6: c.number_format = usdm
        r = r + 13
        ws.cell(r, 1, "ETF flows (net buys)").font = B; ws.cell(r, 9, "ETF flows (net sells)").font = B
        for i in range(10):
            for off, side in ((0, "etf_buys"), (8, "etf_sells")):
                if i < len(t[side]):
                    x = t[side][i]
                    ws.cell(r + 1 + i, off + 2, x["issuer"]); ws.cell(r + 1 + i, off + 3, x["flow"]).number_format = usdm
        widths(ws, [4, 34, 11, 8, 8, 7, 10, 3, 4, 34, 11, 8, 8, 7, 10])
        ws.freeze_panes = "A5"

    # Top 10 managers
    ws = wb.create_sheet("Top 10 managers")
    header(ws, 1, ["Manager", "AUM", "AUM QoQ", "Net flow", "Turnover", "Top 5 buys", "Top 5 sells"])
    for i, mg in enumerate(latest["managers"], 2):
        ws.cell(i, 1, mg["name"]); ws.cell(i, 2, mg["aum"]).number_format = usd
        ws.cell(i, 3, mg["aum"] / mg["aum_prior"] - 1 if mg["aum_prior"] else None).number_format = pct
        ws.cell(i, 4, mg["net_flow"]).number_format = usd; ws.cell(i, 5, mg["turnover"]).number_format = pct
        ws.cell(i, 6, "; ".join(f"{b['issuer']} {b['flow']/1e9:+.1f}bn" for b in mg["buys"]))
        ws.cell(i, 7, "; ".join(f"{b['issuer']} {b['flow']/1e9:+.1f}bn" for b in mg["sells"]))
    widths(ws, [34, 12, 10, 12, 10, 80, 80])

    # Trends
    ws = wb.create_sheet("Trends")
    qs = sorted(hist["quarters"])
    ws["A1"] = "Quarter-by-quarter trends by tier"; ws["A1"].font = T
    r = 3
    for metric, fmt, lab in (("net_flow_pct", pct, "Net flow % AUM"), ("turnover", pct, "Turnover"),
                             ("stock_net_flow", usd, "Stock net flow"), ("etf_net_flow", usd, "ETF net flow"),
                             ("put_call", "0.00", "Put/Call notional"), ("median_top10_conc", pct, "Median top-10 concentration")):
        ws.cell(r, 1, lab).font = B
        header(ws, r + 1, ["Tier"] + qs)
        for i, (tk, tl, _) in enumerate(TIERS + [("all", "All", None)]):
            ws.cell(r + 2 + i, 1, hist["quarters"][qs[-1]]["tiers"][tk]["label"])
            for j, qq in enumerate(qs, 2):
                v = hist["quarters"][qq]["tiers"].get(tk, {}).get(metric)
                ws.cell(r + 2 + i, j, v).number_format = fmt
        r += 9
    # persistent names
    ws.cell(r, 1, "Persistent flows: stocks in a tier's top 20 for 2+ quarters").font = B; r += 1
    header(ws, r, ["Tier", "Issuer", "Side", "Quarters", "Cumulative flow"]); r += 1
    for tk, tl, _ in TIERS:
        for side in ("top_buys", "top_sells"):
            seen = defaultdict(list)
            for qq in qs:
                for x in hist["quarters"][qq]["tiers"].get(tk, {}).get(side, []):
                    seen[x["issuer"]].append((qq, x["flow"]))
            for n, v in sorted(seen.items(), key=lambda kv: -abs(sum(f for _, f in kv[1]))):
                if len(v) >= 2:
                    ws.cell(r, 1, tl); ws.cell(r, 2, n); ws.cell(r, 3, "BUY" if side == "top_buys" else "SELL")
                    ws.cell(r, 4, ", ".join(qq for qq, _ in v)); ws.cell(r, 5, sum(f for _, f in v)).number_format = usd; r += 1
    widths(ws, [28, 34, 10, 30, 14, 14, 14, 14])

    # Method
    ws = wb.create_sheet("Method")
    notes = [
        "Source: SEC Form 13F bulk data sets (all 13F-HR filers). Latest restated/original filing per manager-quarter, plus NEW HOLDINGS amendments.",
        "Tiers by current-quarter long-equity 13F value: Top 10 by size; $10B+ ex top 10; $1-10B; $100M-1B. Sub-$100M voluntary filers excluded.",
        "Flow = (shares this Q - shares last Q, split-adjusted) x median implied price this Q. This shows active buying and selling, not price drift.",
        "Only managers filing in both quarters are counted. New and dropped filers are reported in the universe counts.",
        "Long common/ETF positions only (SH). Options appear separately as put/call notional, which is the underlying value, not the premium.",
        "Turnover = min(gross buys, gross sells) / average AUM, which approximates repositioning. Gross activity = (buys + sells) / average AUM.",
        "Caveats: lagging (45 days after quarter end), long-only, US-listed 13F securities only, no shorts or cash. Large passive managers' flows reflect fund subscriptions and redemptions.",
        "Splits are detected when the median share ratio across 5+ continuing holders is close to a standard split ratio.",
    ]
    for i, n in enumerate(notes, 1):
        ws.cell(i, 1, n)
    ws.column_dimensions["A"].width = 160
    wb.save(path)


# ----------------------------------------------------------------------------- main
def zips_for_quarters(urls, quarters):
    """Each quarter's filings land mostly in the window starting 2 months after quarter-end-month... pick all windows
    covering the period from quarter end to +5 months (captures late filers)."""
    def window(u):
        m = re.search(r"(\d{2}[a-z]{3}\d{4})-(\d{2}[a-z]{3}\d{4})_form13f", u)
        return dt.datetime.strptime(m.group(1), "%d%b%Y"), dt.datetime.strptime(m.group(2), "%d%b%Y")
    need = set()
    for q in quarters:
        y, n = int(q[:4]), int(q[-1])
        qend = dt.datetime(y, 3 * n, 1) + pd.offsets.MonthEnd(0)
        lo, hi = qend, qend + pd.DateOffset(months=5)
        for u in urls:
            s, e = window(u)
            if e >= lo and s <= hi:
                need.add(u)
    return sorted(need, key=lambda u: window(u)[0])


def latest_complete_quarter(urls):
    """Latest quarter whose main filing window (quarter end + 2 months) is fully covered by a posted data set."""
    last_end = max(dt.datetime.strptime(re.search(r"-(\d{2}[a-z]{3}\d{4})_form13f", u).group(1), "%d%b%Y") for u in urls)
    d = last_end - pd.DateOffset(months=2)
    d = d.to_pydatetime()
    q = (d.month - 1) // 3 + 1
    qend_month = 3 * q
    if d.month != qend_month:  # step back to last fully-ended quarter
        q -= 1
        if q == 0:
            return f"{d.year - 1}Q4"
    return f"{d.year}Q{q}"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--quarters", type=int, default=1)
    ap.add_argument("--history", default=None)
    ap.add_argument("--target", default=None, help="e.g. 2026Q2; default = latest complete")
    a = ap.parse_args()
    os.makedirs(OUT, exist_ok=True)

    urls = list_zips()
    tgt = a.target or latest_complete_quarter(urls)
    qs = [tgt]
    for _ in range(a.quarters - 1):
        qs.insert(0, prev_q(qs[0]))
    hist = {"quarters": {}}
    if a.history and os.path.exists(a.history):
        hist = json.load(open(a.history))
    todo = [q for q in qs if q not in hist["quarters"]] or [tgt]
    need_q = sorted(set(todo + [prev_q(q) for q in todo]))
    paths = [fetch(u) for u in zips_for_quarters(urls, need_q)]
    print("using", [os.path.basename(p) for p in paths], file=sys.stderr)
    info, names = build_books(paths)
    info = fix_units(info)
    for q in todo:
        print("computing", q, file=sys.stderr)
        res, *_ = compute_quarter(info, names, q)
        hist["quarters"][q] = res
    hist["latest"] = max(hist["quarters"])
    hist["updated"] = dt.date.today().isoformat()
    jp = os.path.join(OUT, "history.json")
    json.dump(hist, open(jp, "w"), separators=(",", ":"))
    latest = hist["quarters"][hist["latest"]]
    xp = os.path.join(OUT, f"13F_flows_{hist['latest']}.xlsx")
    write_excel(hist, latest, xp)
    print(json.dumps({"history": jp, "excel": xp, "latest": hist["latest"], "quarters": sorted(hist["quarters"])}))


if __name__ == "__main__":
    main()
